"""analysis.py - EDA, kiểm tra pipeline, phân tích lỗi, điểm thưởng, biểu đồ và results.xlsx.

Gọi từ notebook:
    eda(images_dir, labels_dir, out)          # Bước 0: phân bố lớp, ảnh mẫu, thống kê ảnh, kiểm tra split
    sanity_checks(lab)                        # Bước 0: loss ban đầu ~ ln 9, overfit 1 batch, ảnh sau aug
    robustness(lab)                           # thưởng: lệch phân phối (tối/nhiễu/mờ), ECE trước/sau TS
    error_analysis(lab)                       # ma trận nhầm lẫn, ảnh đoán sai, Grad-CAM (sau chung kết)
    make_plots(lab)                           # đánh đổi F1-độ trễ, ablation, reliability
    build_workbook(lab, path)                 # results.xlsx (7 sheet GUIDE mục 6.1 + phụ)
    export(lab, dest)                         # gom file nhỏ để nộp (không checkpoint)
"""
from __future__ import annotations

import json
import math
import shutil
from pathlib import Path

import numpy as np
import pandas as pd

import dataset as D

PAPER_COUNTS = {"Chinee Apple": 1125, "Lantana": 1064, "Parkinsonia": 1031, "Parthenium": 1022,
                "Prickly Acacia": 1062, "Rubber Vine": 1009, "Siam Weed": 1074, "Snake Weed": 1016,
                "Negatives": 9106}


def _plt():
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    return plt


def _softmax(z):
    z = np.asarray(z, dtype=np.float64)
    z = z - z.max(1, keepdims=True)
    e = np.exp(z)
    return e / e.sum(1, keepdims=True)


# --------------------------------------------------------------------------- #
# Bước 0: EDA
# --------------------------------------------------------------------------- #
def eda(images_dir, labels_dir, out, n_stats: int = 2000) -> dict:
    from PIL import Image

    out = Path(out)
    fig_dir = out / "figures"
    fig_dir.mkdir(parents=True, exist_ok=True)
    tr, va, te = D.load_split(labels_dir, 0)
    rep = D.check_split(tr, va, te, images_dir, Path(labels_dir) / "labels.csv")
    pc = pd.DataFrame(rep["per_class"])
    pc["paper_table1"] = [PAPER_COUNTS[c] for c in pc.index]
    pc["diff_vs_paper"] = pc["all"] - pc["paper_table1"]
    pc["pct_all"] = pc["all"] / pc["all"].sum() * 100
    pc.to_csv(out / "eda_class_counts.csv")
    rep["max_min_ratio"] = float(pc["all"].max() / pc["all"].min())
    rep["max_min_ratio_train"] = float(pc["train"].max() / pc["train"].min())

    plt = _plt()
    fig, ax = plt.subplots(1, 2, figsize=(15, 4.5))
    x = np.arange(len(pc))
    for i, s in enumerate(("train", "val", "test")):
        ax[0].bar(x + (i - 1) * 0.27, pc[s], 0.27, label=s)
    ax[0].set_xticks(x, pc.index, rotation=35, ha="right")
    ax[0].set(ylabel="số ảnh", title="Phân bố lớp theo tập (fold 0)")
    ax[0].legend()
    ax[1].bar(x - 0.2, pc["all"], 0.4, label="đếm thật (train+val+test)")
    ax[1].bar(x + 0.2, pc["paper_table1"], 0.4, label="Table 1 bài báo")
    ax[1].set_yscale("log")
    ax[1].set_xticks(x, pc.index, rotation=35, ha="right")
    ax[1].set(ylabel="số ảnh (log)", title=f"Đối chiếu Table 1; lớn nhất/nhỏ nhất = {rep['max_min_ratio']:.2f}")
    ax[1].legend()
    for a in ax:
        a.grid(alpha=.3, axis="y")
    fig.tight_layout()
    fig.savefig(fig_dir / "eda_class_distribution.png", dpi=130)
    plt.close(fig)

    # ảnh mẫu: 4 ảnh mỗi lớp (train)
    rng = np.random.default_rng(0)
    fig, axes = plt.subplots(9, 4, figsize=(8, 18))
    for c in range(9):
        files = tr.loc[tr["Label"] == c, "Filename"].to_numpy()
        for j, f in enumerate(rng.choice(files, 4, replace=False)):
            axes[c, j].imshow(Image.open(Path(images_dir) / f))
            axes[c, j].set_xticks([]), axes[c, j].set_yticks([])
            if j == 0:
                axes[c, j].set_ylabel(D.CLASS_NAMES[c], fontsize=9)
    fig.suptitle("Ảnh mẫu theo lớp (train, ngẫu nhiên seed 0)")
    fig.tight_layout()
    fig.savefig(fig_dir / "eda_samples.png", dpi=110)
    plt.close(fig)

    # thống kê ảnh: kích thước mọi ảnh (đọc header), mean/std kênh trên n_stats ảnh train
    sizes, modes = {}, {}
    for f in pd.concat([tr, va, te])["Filename"]:
        with Image.open(Path(images_dir) / f) as im:
            sizes[im.size] = sizes.get(im.size, 0) + 1
            modes[im.mode] = modes.get(im.mode, 0) + 1
    sample = rng.choice(tr["Filename"].to_numpy(), min(n_stats, len(tr)), replace=False)
    acc = np.zeros(3)
    acc2 = np.zeros(3)
    for f in sample:
        a = np.asarray(Image.open(Path(images_dir) / f).convert("RGB"), dtype=np.float64) / 255
        acc += a.mean((0, 1))
        acc2 += (a ** 2).mean((0, 1))
    mean = acc / len(sample)
    std = np.sqrt(acc2 / len(sample) - mean ** 2)
    rep["image_sizes"] = {f"{k[0]}x{k[1]}": v for k, v in sizes.items()}
    rep["image_modes"] = modes
    rep["train_channel_mean"] = mean.round(4).tolist()
    rep["train_channel_std"] = std.round(4).tolist()
    rep["imagenet_mean"], rep["imagenet_std"] = list(D.IMAGENET_MEAN), list(D.IMAGENET_STD)
    (out / "eda.json").write_text(json.dumps(rep, indent=2, ensure_ascii=False))
    print(f"Kích thước ảnh: {rep['image_sizes']}; mode {modes}")
    print(f"Mean/std kênh (train, {len(sample)} ảnh): {rep['train_channel_mean']} / {rep['train_channel_std']}")
    print(f"Tỉ lệ lớp lớn nhất / nhỏ nhất: {rep['max_min_ratio']:.2f}")
    return rep


# --------------------------------------------------------------------------- #
# Bước 0: kiểm tra pipeline (slide trang 59, GUIDE mục 1.3)
# --------------------------------------------------------------------------- #
def sanity_checks(lab, backbone: str = "resnet50", n_overfit: int = 16, steps: int = 120) -> dict:
    import torch
    import torch.nn as nn

    import train as TR
    from losses import FocalLoss, mix_batch
    from model import build_model, param_groups, set_train_mode, weight_info

    out = lab.out / "sanity"
    out.mkdir(exist_ok=True)
    path = out / "sanity.json"
    if path.exists():
        print("sanity đã chạy:", json.loads(path.read_text()))
        return json.loads(path.read_text())
    plt = _plt()
    TR.set_seed(0)
    dev = torch.device("cuda" if torch.cuda.is_available() else "cpu")
    tr, va, _ = D.load_split(lab.labels_dir, 0)
    model = build_model(backbone, pretrained=True).to(dev)
    wi = weight_info(model)
    mean, std = tuple(wi["mean"]), tuple(wi["std"])
    res = {"backbone": wi["timm_name"], "expected_initial_loss": math.log(9)}

    # 1) loss ban đầu với head mới (eval mode, 4 batch val)
    ld = D.make_loader(va.sample(256, random_state=0), lab.images_dir, D.build_transforms(False, 224, mean=mean, std=std),
                       64, train=False, num_workers=2, cache_path=lab.cache_path)
    ce = nn.CrossEntropyLoss()
    with torch.inference_mode():
        model.eval()
        losses, ent = [], []
        for x, y, _ in ld:
            z = model(x.to(dev)).float()
            losses.append(ce(z, y.to(dev)).item())
            ent.append(float(-(z.softmax(1) * z.log_softmax(1)).sum(1).mean()))
    res["initial_loss"] = float(np.mean(losses))
    res["initial_pred_entropy"] = float(np.mean(ent))
    # focal gamma=0 == CE trên logit thật
    zz = torch.randn(64, 9)
    yy = torch.randint(0, 9, (64,))
    res["focal_gamma0_minus_ce"] = float(abs(FocalLoss(0.0)(zz, yy) - ce(zz, yy)))

    # 2) overfit một batch nhỏ (16 ảnh, không augmentation)
    sub = tr.groupby("Label").sample(2, random_state=0).iloc[:n_overfit]
    ds = D.DeepWeedsDataset(sub, lab.images_dir, D.build_transforms(False, 224, mean=mean, std=std), lab.cache_path)
    xb = torch.stack([ds[i][0] for i in range(len(ds))]).to(dev)
    yb = torch.tensor([ds[i][1] for i in range(len(ds))], device=dev)
    opt = torch.optim.AdamW(param_groups(model, 1e-4, 1e-3, 0.0))
    curve = []
    for _ in range(steps):
        set_train_mode(model)
        loss = ce(model(xb).float(), yb)
        opt.zero_grad()
        loss.backward()
        opt.step()
        curve.append(loss.item())
    res["overfit_final_loss"] = curve[-1]
    res["overfit_ok"] = bool(curve[-1] < 0.05)
    fig, ax = plt.subplots(figsize=(6, 3.5))
    ax.plot(curve)
    ax.axhline(math.log(9), ls="--", c="gray", label="ln 9")
    ax.set_yscale("log")
    ax.set(xlabel="bước", ylabel="loss CE", title=f"Overfit {len(sub)} ảnh ({backbone}): loss cuối {curve[-1]:.4f}")
    ax.legend()
    fig.tight_layout()
    fig.savefig(out / "overfit_one_batch.png", dpi=120)
    plt.close(fig)

    # 3) ảnh sau augmentation (đã giải chuẩn hoá) + nhãn
    show = tr.sample(6, random_state=1)
    rows = list(D.AUG_CHOICES) + ["mixup", "cutmix"]
    fig, axes = plt.subplots(len(rows), 6, figsize=(13, 2.3 * len(rows)))
    for r, aug in enumerate(rows):
        tf = D.build_transforms(True, 224, aug if aug in D.AUG_CHOICES else "basic", mean, std)
        dsa = D.DeepWeedsDataset(show, lab.images_dir, tf, lab.cache_path)
        xs = torch.stack([dsa[i][0] for i in range(6)])
        ys = torch.tensor([dsa[i][1] for i in range(6)])
        title = [D.CLASS_NAMES[int(v)] for v in ys]
        if aug in ("mixup", "cutmix"):
            torch.manual_seed(3)
            xs, (ya, yb2, lam) = mix_batch(xs, ys, 1.0 if aug == "cutmix" else 0.4, aug,
                                           np.random.default_rng(3))
            title = [f"{D.CLASS_NAMES[int(a)][:10]}/{D.CLASS_NAMES[int(b)][:10]}\nλ={lam:.2f}" for a, b in zip(ya, yb2)]
        imgs = D.denormalize(xs, mean, std).permute(0, 2, 3, 1).numpy()
        for j in range(6):
            axes[r, j].imshow(imgs[j])
            axes[r, j].set_title(title[j], fontsize=7)
            axes[r, j].axis("off")
        axes[r, 0].text(-0.15, 0.5, aug, transform=axes[r, 0].transAxes, rotation=90, va="center", ha="right")
    fig.suptitle("Ảnh sau augmentation (giải chuẩn hoá) và nhãn")
    fig.tight_layout()
    fig.savefig(out / "augmentations.png", dpi=110)
    plt.close(fig)
    (path).write_text(json.dumps(res, indent=2))
    print(json.dumps(res, indent=2))
    del model
    return res


# --------------------------------------------------------------------------- #
# Thưởng: lệch phân phối trên VAL (không dùng test)
# --------------------------------------------------------------------------- #
def _corrupt(x01, kind, sev, gen):
    import torch
    import torch.nn.functional as F
    if kind == "clean":
        return x01
    if kind == "dark":
        return (x01 * sev).clamp(0, 1)
    if kind == "low_contrast":
        m = x01.mean((1, 2, 3), keepdim=True)
        return ((x01 - m) * sev + m).clamp(0, 1)
    if kind == "noise":
        return (x01 + sev * torch.randn(x01.shape, generator=gen, device="cpu").to(x01.device)).clamp(0, 1)
    if kind == "blur":
        k = int(2 * round(3 * sev) + 1)
        t = torch.arange(k, device=x01.device, dtype=x01.dtype) - k // 2
        g = torch.exp(-t ** 2 / (2 * sev ** 2))
        g = g / g.sum()
        c = x01.shape[1]
        x = F.conv2d(F.pad(x01, (k // 2,) * 4, mode="reflect"), g.view(1, 1, 1, k).repeat(c, 1, 1, 1), groups=c)
        return F.conv2d(x, g.view(1, 1, k, 1).repeat(c, 1, 1, 1), groups=c)
    raise ValueError(kind)


def robustness(lab, exp_id: str = "F01", seed: int = 0) -> pd.DataFrame:
    import torch

    import inference as INF
    import train as TR
    from eval import compute_metrics

    path = lab.out / "inference" / "robustness.csv"
    if path.exists():
        return pd.read_csv(path)
    dev = torch.device("cuda" if torch.cuda.is_available() else "cpu")
    cfg = TR.load_config(exp_id, seed, str(lab.out / "runs"))
    model, mean, std = TR.load_trained(cfg, dev)
    _, va, _ = D.load_split(cfg.labels_dir, 0)
    if cfg.eval_limit:
        va = va.iloc[: cfg.eval_limit].reset_index(drop=True)
    loader = TR.build_eval_loader(cfg, va, mean, std, batch_size=128)
    m_t = torch.tensor(mean, device=dev).view(1, 3, 1, 1)
    s_t = torch.tensor(std, device=dev).view(1, 3, 1, 1)
    settings = [("clean", 0), ("dark", 0.5), ("dark", 0.3), ("low_contrast", 0.5), ("noise", 0.05),
                ("noise", 0.1), ("blur", 1.0), ("blur", 2.0)]
    logits = {}
    for kind, sev in settings:
        gen = torch.Generator().manual_seed(0)
        ys, zs = [], []
        model.eval()
        with torch.inference_mode():
            for x, y, _ in loader:
                x01 = (x.to(dev) * s_t + m_t).clamp(0, 1)
                xc = (_corrupt(x01, kind, sev, gen) - m_t) / s_t
                zs.append(model(xc).float().cpu())
                ys.append(y)
        logits[(kind, sev)] = torch.cat(zs).numpy()
        y = torch.cat(ys).numpy()
    T = INF.fit_temperature(logits[("clean", 0)], y)
    rows = []
    for (kind, sev), z in logits.items():
        p0, p1 = _softmax(z), INF.apply_temperature(z, T)
        m0 = compute_metrics(y, p0.argmax(1), p0)
        m1 = compute_metrics(y, p1.argmax(1), p1)
        rows.append({"corruption": kind, "severity": sev, "val_macro_f1": m0["macro_f1"], "val_top1": m0["top1"],
                     "ece_uncal": m0["ece"], "ece_ts_clean_T": m1["ece"], "T_clean_val": T,
                     "f1_chinee_apple": m0["f1"][0], "f1_snake_weed": m0["f1"][7]})
    df = pd.DataFrame(rows)
    df.to_csv(path, index=False)
    print(df.round(4).to_string())
    del model
    return df


# --------------------------------------------------------------------------- #
# Sau chung kết: phân tích lỗi trên test
# --------------------------------------------------------------------------- #
def _gradcam(model, x, cls):
    """Grad-CAM tổng quát cho timm: lấy gradient tại đầu ra forward_features (CNN: B,C,H,W;
    Swin: B,H,W,C; ViT: B,N,D -> bỏ token tiền tố, xếp lại thành lưới)."""
    import torch
    import torch.nn.functional as F
    model.eval()
    x = x.clone().requires_grad_(True)
    feats = model.forward_features(x)
    feats.retain_grad()
    logits = model.forward_head(feats)
    logits[torch.arange(len(cls)), cls].sum().backward()
    f, g = feats.detach(), feats.grad.detach()
    nf = getattr(model, "num_features", None)
    if f.ndim == 3:  # ViT tokens
        npre = getattr(model, "num_prefix_tokens", 1)
        f, g = f[:, npre:], g[:, npre:]
        side = int(round(math.sqrt(f.shape[1])))
        f = f.transpose(1, 2).reshape(f.shape[0], -1, side, side)
        g = g.transpose(1, 2).reshape(g.shape[0], -1, side, side)
    elif f.ndim == 4 and f.shape[1] != nf and f.shape[-1] == nf:  # NHWC
        f, g = f.permute(0, 3, 1, 2), g.permute(0, 3, 1, 2)
    w = g.mean((2, 3), keepdim=True)
    cam = F.relu((w * f).sum(1, keepdim=True))
    cam = F.interpolate(cam, size=x.shape[-2:], mode="bilinear", align_corners=False)[:, 0]
    cam = cam / cam.flatten(1).max(1)[0].clamp(min=1e-8).view(-1, 1, 1)
    return cam.cpu().numpy(), logits.detach().float().softmax(1).cpu().numpy()


def plot_confusion(cm, title, path, normalize=True):
    plt = _plt()
    cmn = cm / cm.sum(1, keepdims=True).clip(min=1) if normalize else cm
    fig, ax = plt.subplots(figsize=(8.5, 7.5))
    im = ax.imshow(cmn, cmap="Blues", vmin=0, vmax=1 if normalize else None)
    for i in range(cm.shape[0]):
        for j in range(cm.shape[1]):
            ax.text(j, i, f"{cm[i, j]}\n{cmn[i, j] * 100:.1f}%" if normalize else f"{cm[i, j]}",
                    ha="center", va="center", fontsize=7, color="white" if cmn[i, j] > .5 else "black")
    ax.set_xticks(range(9), D.CLASS_NAMES, rotation=40, ha="right")
    ax.set_yticks(range(9), D.CLASS_NAMES)
    ax.set(xlabel="dự đoán", ylabel="nhãn thật", title=title)
    fig.colorbar(im, ax=ax, fraction=.046)
    fig.tight_layout()
    fig.savefig(path, dpi=130)
    plt.close(fig)


def error_analysis(lab, exp_id: str = "F01", seed: int = 0, n_show: int = 12) -> dict:
    import torch
    from PIL import Image

    import train as TR
    from eval import confusion_matrix

    fig_dir = lab.out / "figures"
    plt = _plt()
    pdir = lab.out / "predictions"
    info = {}
    for tag in ("F01", "T00", "F02"):
        files = sorted(pdir.glob(f"{tag}_seed*_test.csv"))
        if not files:
            continue
        cm = sum(confusion_matrix(pd.read_csv(f)["y_true"].to_numpy(), pd.read_csv(f)["y_pred"].to_numpy())
                 for f in files)
        plot_confusion(cm, f"{tag}: ma trận nhầm lẫn test (cộng {len(files)} seed, % theo hàng)",
                       fig_dir / f"confusion_{tag}_test.png")
        pd.DataFrame(cm, index=D.CLASS_NAMES, columns=D.CLASS_NAMES).to_csv(lab.out / "eval_out" / f"{tag}_confusion_test.csv")
        info[tag] = {"chinee_to_snake": int(cm[0, 7]), "snake_to_chinee": int(cm[7, 0]),
                     "parkinsonia_to_prickly": int(cm[2, 4]), "total_errors": int(cm.sum() - np.trace(cm))}
    pf = pd.read_csv(pdir / f"{exp_id}_seed{seed}_test.csv")
    wrong = pf[pf["y_true"] != pf["y_pred"]].copy()
    probs = pf[[f"p{i}" for i in range(9)]].to_numpy()
    wrong["conf"] = probs[wrong.index, wrong["y_pred"].to_numpy()]
    hard = wrong[wrong["y_true"].isin([0, 7]) & wrong["y_pred"].isin([0, 7])].sort_values("conf", ascending=False)
    other = wrong.drop(hard.index).sort_values("conf", ascending=False)
    pick = pd.concat([hard.head(n_show // 2), other.head(n_show - min(len(hard), n_show // 2))]).head(n_show)
    info["n_wrong_seed0"] = int(len(wrong))
    info["n_chinee_snake_confusions_seed0"] = int(len(hard))
    if len(pick):
        dev = torch.device("cuda" if torch.cuda.is_available() else "cpu")
        cfg = TR.load_config(exp_id, seed, str(lab.out / "runs"))
        model, mean, std = TR.load_trained(cfg, dev)
        tf = D.build_transforms(False, cfg.img_size, mean=mean, std=std)
        ds = D.DeepWeedsDataset(pick, lab.images_dir, tf, lab.cache_path)
        x = torch.stack([ds[i][0] for i in range(len(ds))]).to(dev)
        try:
            cam_pred, _ = _gradcam(model, x, torch.tensor(pick["y_pred"].to_numpy(), device=dev))
            cam_true, _ = _gradcam(model, x, torch.tensor(pick["y_true"].to_numpy(), device=dev))
        except Exception as e:  # noqa: BLE001
            print("Grad-CAM lỗi:", e)
            cam_pred = cam_true = None
        imgs = D.denormalize(x.detach(), mean, std).permute(0, 2, 3, 1).cpu().numpy()
        ncol = 3 if cam_pred is not None else 1
        fig, axes = plt.subplots(len(pick), ncol, figsize=(3.2 * ncol, 3.1 * len(pick)))
        axes = np.atleast_2d(axes).reshape(len(pick), ncol)
        for i, (_, r) in enumerate(pick.iterrows()):
            t, p = D.CLASS_NAMES[int(r["y_true"])], D.CLASS_NAMES[int(r["y_pred"])]
            axes[i, 0].imshow(imgs[i])
            axes[i, 0].set_title(f"{r['Filename']}\nthật: {t} | đoán: {p} ({r['conf']:.2f})", fontsize=7)
            if cam_pred is not None:
                for j, (cam, lbl) in enumerate(((cam_pred, f"Grad-CAM lớp đoán: {p}"), (cam_true, f"Grad-CAM lớp thật: {t}")), 1):
                    axes[i, j].imshow(imgs[i])
                    axes[i, j].imshow(cam[i], cmap="jet", alpha=.45)
                    axes[i, j].set_title(lbl, fontsize=7)
            for a in axes[i]:
                a.axis("off")
        fig.suptitle(f"{exp_id} seed{seed}: ảnh test bị đoán sai (ưu tiên Chinee apple <-> Snake weed)", fontsize=9)
        fig.tight_layout()
        fig.savefig(fig_dir / f"errors_gradcam_{exp_id}_seed{seed}.png", dpi=100)
        plt.close(fig)
        del model
    # ảnh gốc độ phân giải đầy đủ của các cặp nhầm Chinee <-> Snake để so bằng mắt
    if len(hard):
        k = min(8, len(hard))
        fig, axes = plt.subplots(2, k, figsize=(2.6 * k, 5.6), squeeze=False)
        for j, (_, r) in enumerate(hard.head(k).iterrows()):
            axes[0, j].imshow(Image.open(Path(lab.images_dir) / r["Filename"]))
            axes[0, j].set_title(f"thật {D.CLASS_NAMES[int(r['y_true'])]}\nđoán {D.CLASS_NAMES[int(r['y_pred'])]}", fontsize=7)
            ref_cls = int(r["y_pred"])
            tr, _, _ = D.load_split(lab.labels_dir, 0)
            ref = tr[tr["Label"] == ref_cls].sample(1, random_state=j)["Filename"].iloc[0]
            axes[1, j].imshow(Image.open(Path(lab.images_dir) / ref))
            axes[1, j].set_title(f"ảnh train lớp {D.CLASS_NAMES[ref_cls]}", fontsize=7)
            for a in axes[:, j]:
                a.axis("off")
        fig.suptitle("Hàng trên: ảnh test bị nhầm; hàng dưới: một ảnh train của lớp bị đoán thành", fontsize=9)
        fig.tight_layout()
        fig.savefig(fig_dir / "chinee_snake_confusions.png", dpi=100)
        plt.close(fig)
    (lab.out / "eval_out" / "error_analysis.json").write_text(json.dumps(info, indent=2, ensure_ascii=False))
    print(json.dumps(info, indent=2, ensure_ascii=False))
    return info


# --------------------------------------------------------------------------- #
# Biểu đồ tổng hợp
# --------------------------------------------------------------------------- #
def make_plots(lab) -> None:
    plt = _plt()
    fig_dir = lab.out / "figures"
    bt = lab.backbone_table()
    if len(bt):
        fig, ax = plt.subplots(1, 2, figsize=(13, 4.8))
        for _, r in bt.iterrows():
            for a, xk in ((ax[0], "latency_b1_p95_ms"), (ax[1], "gmacs")):
                a.scatter(r[xk], r["val_macro_f1"], s=30 + 6 * r["params_M"], alpha=.6)
                a.annotate(f"{r['exp_id']} {r['backbone']}", (r[xk], r["val_macro_f1"]), fontsize=8,
                           xytext=(4, 4), textcoords="offset points")
        ax[0].set(xlabel="độ trễ p95 batch 1, FP32 (ms) — đo sơ bộ", ylabel="macro-F1 val",
                  title="Backbone: macro-F1 val vs độ trễ (cỡ điểm ~ #params)")
        ax[1].set(xlabel="GMAC / ảnh", ylabel="macro-F1 val", title="Backbone: macro-F1 val vs GMAC")
        for a in ax:
            a.grid(alpha=.3)
        fig.tight_layout()
        fig.savefig(fig_dir / "backbones_tradeoff.png", dpi=130)
        plt.close(fig)
    tt = lab.training_table() if (lab.out / "decisions.json").exists() and "backbone" in lab.decisions else None
    if tt is not None and len(tt):
        st = lab.t00_stats()
        t = tt[tt["exp_id"] != "T00"]
        fig, ax = plt.subplots(figsize=(11, 4.8))
        cols = ["tab:green" if d > st["noise"] else ("tab:red" if d < -st["noise"] else "tab:gray")
                for d in t["delta_vs_T00_mean"]]
        ax.bar(t["exp_id"] + "\n" + t["desc"], t["delta_vs_T00_mean"], color=cols)
        ax.axhspan(-st["noise"], st["noise"], color="gray", alpha=.2, label=f"± std T00 ({st['std']:.4f}, 3 seed)")
        ax.axhline(0, c="k", lw=.8)
        ax.set(ylabel="Δ macro-F1 val so với mean T00", title=f"Ablation công thức huấn luyện (1 seed mỗi dòng); "
                                                              f"T00 = {st['mean']:.4f} ± {st['std']:.4f}")
        ax.tick_params(axis="x", labelsize=7)
        ax.legend()
        ax.grid(alpha=.3, axis="y")
        fig.tight_layout()
        fig.savefig(fig_dir / "training_ablation.png", dpi=130)
        plt.close(fig)
    inf = lab.out / "inference" / "inference.csv"
    if inf.exists():
        df = pd.read_csv(inf).dropna(subset=["val_macro_f1", "lat_p50_ms"])
        fig, ax = plt.subplots(figsize=(10, 5.5))
        ax.scatter(df["lat_p50_ms"], df["val_macro_f1"], c="tab:blue")
        for _, r in df.iterrows():
            ax.annotate(r["exp_id"], (r["lat_p50_ms"], r["val_macro_f1"]), fontsize=8, xytext=(3, 3),
                        textcoords="offset points")
        ax.axvline(100, ls="--", c="tab:red", label="ngân sách 100 ms")
        ax.set_xscale("log")
        ax.set(xlabel="độ trễ p50 batch 1 (ms, log)", ylabel="macro-F1 val",
               title="Đánh đổi độ chính xác - độ trễ của các phương pháp suy luận")
        ax.grid(alpha=.3, which="both")
        ax.legend()
        fig.tight_layout()
        fig.savefig(fig_dir / "inference_tradeoff.png", dpi=130)
        plt.close(fig)
    # reliability diagram test: F01 trước / sau temperature scaling (seed 0)
    p_cal = lab.out / "predictions" / "F01_seed0_test.csv"
    p_unc = lab.out / "predictions" / "F01_uncal_seed0_test.csv"
    if p_cal.exists() and p_unc.exists():
        fig, ax = plt.subplots(figsize=(5.5, 5))
        for p, lbl in ((p_unc, "trước TS"), (p_cal, "sau TS")):
            d = pd.read_csv(p)
            pr = d[[f"p{i}" for i in range(9)]].to_numpy()
            conf, corr = pr.max(1), (pr.argmax(1) == d["y_true"].to_numpy())
            idx = np.clip(np.ceil(conf * 15).astype(int) - 1, 0, 14)
            xs = [conf[idx == m].mean() for m in range(15) if (idx == m).any()]
            ys = [corr[idx == m].mean() for m in range(15) if (idx == m).any()]
            ax.plot(xs, ys, "-o", ms=3, label=lbl)
        ax.plot([0, 1], [0, 1], "--", c="gray")
        ax.set(xlabel="độ tin cậy", ylabel="accuracy", title="Reliability F01 seed0 (test, 15 bin)")
        ax.legend()
        ax.grid(alpha=.3)
        fig.tight_layout()
        fig.savefig(fig_dir / "reliability_F01_test.png", dpi=130)
        plt.close(fig)


# --------------------------------------------------------------------------- #
# results.xlsx
# --------------------------------------------------------------------------- #
def _per_seed_val(pdir: Path, tag: str) -> dict:
    from eval import compute_metrics, read_pred
    out = {}
    for f in sorted(pdir.glob(f"{tag}_seed*_val.csv")):
        p = read_pred(str(f))
        out[p.seed] = compute_metrics(p.y_true, p.y_pred, p.probs)["macro_f1"]
    return out


def _ms(v):
    v = [x for x in v if x is not None and not (isinstance(x, float) and math.isnan(x))]
    if not v:
        return None
    if len(v) == 1:
        return f"{v[0]:.4f} (1 seed)"
    return f"{np.mean(v):.4f} ± {np.std(v, ddof=1):.4f}"


def build_workbook(lab, path) -> Path:
    from openpyxl.styles import Alignment, Font, PatternFill

    out = lab.out
    pdir, edir = out / "predictions", out / "eval_out"
    sheets = {}
    dec = lab.decisions

    bt = lab.backbone_table()
    if len(bt):
        b = pd.DataFrame({
            "exp_id": bt["exp_id"], "backbone": bt["backbone"], "timm name": bt["timm_name"],
            "tag trọng số": bt["tag"], "khởi tạo": bt["init"], "#tham số (M)": bt["params_M"],
            "GMAC": bt["gmacs"], "độ phân giải": 224, "epoch": [lab.summary(e)["epochs"] for e in bt["exp_id"]],
            "best epoch": bt["best_epoch"], "seed": 0, "macro-F1 val": bt["val_macro_f1"], "top-1 val": bt["val_top1"],
            "F1 val Chinee apple": bt["val_f1_chinee_apple"], "F1 val Snake weed": bt["val_f1_snake_weed"],
            "thời gian train/epoch (s)": bt["time_train_per_epoch_s"],
            "độ trễ batch-1 p50 (ms)": bt["latency_b1_p50_ms"], "độ trễ batch-1 p95 (ms)": bt["latency_b1_p95_ms"],
            "ghi chú": ["bonus: DINOv2 đóng băng + linear probe" if e == "B08" else
                        ("được chọn cho Bước 2-4" if dec.get("backbone", {}).get("from_exp") == e else "")
                        for e in bt["exp_id"]]})
        sheets["Backbones"] = b
    if "backbone" in dec:
        tt = lab.training_table()
        if len(tt):
            sheets["Training"] = pd.DataFrame({
                "exp_id": tt["exp_id"], "backbone": tt["backbone"], "trục (A-G)": tt["axis"],
                "khác T00 ở điểm nào": tt["diff_vs_T00"], "seed": tt["seed"], "macro-F1 val": tt["val_macro_f1"],
                "top-1 val": tt["val_top1"], "Δ macro-F1 vs mean T00": tt["delta_vs_T00_mean"],
                "std T00 (3 seed)": tt["T00_std"], "|Δ| > std T00": tt["beyond_noise"],
                "F1 val Chinee apple": tt["val_f1_chinee_apple"], "F1 val Snake weed": tt["val_f1_snake_weed"],
                "ECE val": tt["val_ece"], "best epoch": tt["best_epoch"],
                "thời gian train/epoch (s)": tt["time_train_per_epoch_s"],
                "ghi chú": [("kết hợp: " + dec.get("combo", {}).get("note", "")) if e == "T15" else
                            ("EMA: best raw %.4f / best EMA %.4f" % (r["best_raw_val_macro_f1"], r["best_ema_val_macro_f1"])
                             if isinstance(r["best_ema_val_macro_f1"], float) and not math.isnan(r["best_ema_val_macro_f1"]) else "")
                            for e, (_, r) in zip(tt["exp_id"], tt.iterrows())]})
    if (out / "inference" / "inference.csv").exists():
        inf = pd.read_csv(out / "inference" / "inference.csv")
        sheets["Inference"] = inf.rename(columns={
            "method": "phương pháp", "model": "mô hình/checkpoint", "K": "K (view/model)", "space": "gộp",
            "val_macro_f1": "macro-F1 val", "val_top1": "top-1 val", "val_ece": "ECE val", "val_nll": "NLL val",
            "lat_p50_ms": "p50 b1 (ms)", "lat_p95_ms": "p95 b1 (ms)", "lat_p99_ms": "p99 b1 (ms)",
            "throughput_img_s": "thông lượng (ảnh/s)", "rel_cost_vs_I00": "chi phí tương đối vs I00",
            "note": "ghi chú"})
    # Final
    rows = []
    groups = [("F01", "chung kết (offline)"), ("F01_uncal", "chung kết, CHƯA temperature scaling"),
              ("F02", "chung kết, thời gian thực 1-view + TS"), ("T00", "mốc: công thức nền + I00")]
    rec, inf_d = dec.get("recipe", {}), dec.get("inference", {})
    for tag, label in groups:
        f = edir / f"{tag}_per_seed.csv"
        if not f.exists():
            continue
        ps = pd.read_csv(f)
        val = _per_seed_val(pdir, tag)
        conf = (f"{rec.get('backbone')} + {rec.get('exp_id')} {rec.get('overrides')} + "
                f"{inf_d.get('exp_id')} {inf_d.get('method')}" if tag.startswith("F01") else
                f"{rec.get('backbone')} + {rec.get('exp_id')} {rec.get('overrides')} + I00 1-view" if tag == "F02" else
                f"{rec.get('backbone')} + T00 + I00 1-view")
        for _, r in ps.iterrows():
            rows.append({"exp_id": tag, "vai trò": label, "cấu hình": conf, "seed": r["seed"],
                         "macro-F1 val": val.get(int(r["seed"])), "macro-F1 test": r["macro_f1"],
                         "top-1 test": r["top1"], "balanced acc test": r["balanced_acc"], "ECE test": r["ece"],
                         "NLL test": r["nll"]})
        js = json.loads((edir / f"{tag}_summary.json").read_text())
        rows.append({"exp_id": tag, "vai trò": label + " — TỔNG HỢP", "cấu hình": conf,
                     "seed": f"mean ± std ({len(ps)} seed)", "macro-F1 val": _ms(list(val.values())),
                     "macro-F1 test": f"{js['macro_f1']['mean']:.4f} ± {js['macro_f1']['std']:.4f}",
                     "top-1 test": f"{js['top1']['mean']:.4f} ± {js['top1']['std']:.4f}",
                     "balanced acc test": f"{js['balanced_acc']['mean']:.4f} ± {js['balanced_acc']['std']:.4f}",
                     "ECE test": f"{js['ece']['mean']:.4f} ± {js['ece']['std']:.4f}",
                     "NLL test": f"{js['nll']['mean']:.4f} ± {js['nll']['std']:.4f}"})
    if rows:
        sheets["Final"] = pd.DataFrame(rows)
    pcs = []
    for tag, label in (("F01", "chung kết"), ("T00", "mốc"), ("F02", "thời gian thực")):
        f = edir / f"{tag}_per_class.csv"
        if f.exists():
            d = pd.read_csv(f)
            d.insert(0, "cấu hình", f"{tag} ({label})")
            pcs.append(d.rename(columns={"class": "lớp", "support": "số ảnh test"}))
    if pcs:
        sheets["PerClass"] = pd.concat(pcs, ignore_index=True)
    if (out / "inference" / "latency.csv").exists():
        lat = pd.read_csv(out / "inference" / "latency.csv")
        fin = json.loads((out / "inference" / "latency_final.json").read_text()) if (out / "inference" / "latency_final.json").exists() else {}
        lat = pd.concat([lat, pd.DataFrame(list(fin.values()))], ignore_index=True)
        keep = ["config", "gpu", "dtype", "batch", "img_size", "bn_fused", "k_views", "k_models", "mode",
                "p50", "p95", "p99", "mean", "std", "n", "warmup", "images_per_s", "torch", "includes_preprocessing"]
        sheets["Latency"] = lat[[c for c in keep if c in lat.columns]].rename(columns={
            "config": "cấu hình", "batch": "batch", "bn_fused": "gộp BN", "p50": "p50 (ms)", "p95": "p95 (ms)",
            "p99": "p99 (ms)", "mean": "mean (ms)", "std": "std (ms)", "n": "số lần đo", "images_per_s": "ảnh/s",
            "includes_preprocessing": "tính tiền xử lý"})
    # Summary: top 10 theo macro-F1 val (B, T, I) + chung kết vs mốc trên test
    cand = []
    for _, r in sheets.get("Backbones", pd.DataFrame()).iterrows():
        cand.append({"exp_id": r["exp_id"], "loại": "backbone", "mô tả": r["backbone"], "macro-F1 val": r["macro-F1 val"],
                     "top-1 val": r["top-1 val"], "p95 b1 (ms)": r["độ trễ batch-1 p95 (ms)"],
                     "chi phí/ghi chú": f"{r['GMAC']:.2f} GMAC, {r['#tham số (M)']:.1f}M"})
    for _, r in sheets.get("Training", pd.DataFrame()).iterrows():
        if r["exp_id"] != "T00":
            cand.append({"exp_id": r["exp_id"], "loại": "huấn luyện", "mô tả": r["khác T00 ở điểm nào"],
                         "macro-F1 val": r["macro-F1 val"], "top-1 val": r["top-1 val"], "p95 b1 (ms)": None,
                         "chi phí/ghi chú": f"Δ vs T00 {r['Δ macro-F1 vs mean T00']:+.4f}"})
    if "Training" in sheets:
        st = lab.t00_stats()
        cand.append({"exp_id": "T00", "loại": "huấn luyện (mốc)", "mô tả": "công thức nền, mean 3 seed",
                     "macro-F1 val": st["mean"], "top-1 val": st["top1_mean"], "p95 b1 (ms)": None,
                     "chi phí/ghi chú": f"std {st['std']:.4f}"})
    for _, r in sheets.get("Inference", pd.DataFrame()).iterrows():
        if pd.notna(r["macro-F1 val"]):
            cand.append({"exp_id": r["exp_id"], "loại": "suy luận", "mô tả": r["phương pháp"],
                         "macro-F1 val": r["macro-F1 val"], "top-1 val": r["top-1 val"], "p95 b1 (ms)": r["p95 b1 (ms)"],
                         "chi phí/ghi chú": (f"x{r['chi phí tương đối vs I00']:.2f} vs I00"
                                             if pd.notna(r["chi phí tương đối vs I00"]) else "")})
    summ = pd.DataFrame(cand)
    if len(summ):
        summ = summ.sort_values("macro-F1 val", ascending=False).head(10).reset_index(drop=True)
        summ.insert(0, "hạng", range(1, len(summ) + 1))
    fin_rows = []
    for tag in ("F01", "F02", "T00"):
        f = edir / f"{tag}_summary.json"
        if f.exists():
            js = json.loads(f.read_text())
            fin_rows.append({"hạng": "", "exp_id": tag, "loại": "TEST (chung kết/mốc)",
                             "mô tả": {"F01": "chung kết offline", "F02": "chung kết thời gian thực", "T00": "mốc T00+I00"}[tag],
                             "macro-F1 val": None,
                             "top-1 val": None, "p95 b1 (ms)": None,
                             "chi phí/ghi chú": f"macro-F1 test {js['macro_f1']['mean']:.4f} ± {js['macro_f1']['std']:.4f}; "
                                                f"top-1 test {js['top1']['mean']:.4f} ± {js['top1']['std']:.4f}; "
                                                f"recall Chinee {js['recall']['mean'][0]:.3f}, Snake {js['recall']['mean'][7]:.3f}"})
    if fin_rows:
        summ = pd.concat([summ, pd.DataFrame([{}]), pd.DataFrame(fin_rows)], ignore_index=True)
    sheets["Summary"] = summ
    if (out / "inference" / "robustness.csv").exists():
        sheets["Robustness"] = pd.read_csv(out / "inference" / "robustness.csv")
    if dec:
        sheets["Decisions"] = pd.DataFrame([{"quyết định": k, "giá trị": json.dumps(
            {kk: vv for kk, vv in v.items() if kk not in ("table", "candidates")}, ensure_ascii=False, default=str)}
            for k, v in dec.items()])

    order = ["Summary", "Backbones", "Training", "Inference", "Final", "PerClass", "Latency", "Robustness", "Decisions"]
    path = Path(path)
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        for name in order:
            if name not in sheets:
                continue
            df = sheets[name]
            df.to_excel(xw, sheet_name=name, index=False)
            ws = xw.sheets[name]
            ws.freeze_panes = "B2"
            for c in ws[1]:
                c.font = Font(bold=True)
                c.alignment = Alignment(wrap_text=True, vertical="top")
            for col in ws.columns:
                width = max(len(str(c.value)) if c.value is not None else 0 for c in col[:60])
                ws.column_dimensions[col[0].column_letter].width = min(48, max(9, width + 2))
                for c in col[1:]:
                    if isinstance(c.value, float):
                        c.number_format = "0.0000"
            hl = PatternFill("solid", fgColor="FFF2CC")
            f1col = next((i for i, c in enumerate(ws[1]) if c.value in ("macro-F1 val", "macro-F1 test")), None)
            if f1col is not None and name in ("Backbones", "Training", "Inference", "Summary"):
                vals = [(r, row[f1col].value) for r, row in enumerate(ws.iter_rows(min_row=2), 2)
                        if isinstance(row[f1col].value, (int, float))]
                if vals:
                    best_r = max(vals, key=lambda t: t[1])[0]
                    for c in ws[best_r]:
                        c.fill = hl
            if name == "Final":
                for row in ws.iter_rows(min_row=2):
                    if "TỔNG HỢP" in str(row[1].value):
                        for c in row:
                            c.fill = hl
                            c.font = Font(bold=True)
    print("Đã ghi", path)
    return path


# --------------------------------------------------------------------------- #
# Gom kết quả để nộp
# --------------------------------------------------------------------------- #
def export(lab, dest) -> Path:
    """Chép file nhỏ (không checkpoint) vào `dest` theo cấu trúc thư mục bài nộp, rồi nén zip."""
    src, dest = lab.out, Path(dest)
    if dest.exists():
        shutil.rmtree(dest)
    dest.mkdir(parents=True)
    # eval_out/ và runs/ nằm trong .gitignore của repo gốc -> đổi tên khi gom để commit được
    for d, t in (("curves", "curves"), ("predictions", "predictions"), ("figures", "figures"),
                 ("eval_out", "eval_results"), ("inference", "inference"), ("sanity", "sanity")):
        if (src / d).exists():
            shutil.copytree(src / d, dest / t)
    for f in ("decisions.json", "eda.json", "eda_class_counts.csv", "results.xlsx"):
        if (src / f).exists():
            shutil.copy2(src / f, dest / f)
    keep = ("config.json", "history.csv", "lr_trace.csv", "summary.json", "val_outputs.npz", "final_predict.json",
            "final_val_views.npz", "final_test_views.npz", "test_outputs.npz")
    for f in (src / "runs").rglob("*"):
        if f.is_file() and f.name in keep:
            t = dest / "logs" / f.relative_to(src / "runs")
            t.parent.mkdir(parents=True, exist_ok=True)
            shutil.copy2(f, t)
    for f in (src / "logs").rglob("worker*.log"):
        t = dest / "logs" / "workers" / f"{f.parent.name}_{f.name}"
        t.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(f, t)
    z = shutil.make_archive(str(dest), "zip", root_dir=dest)
    size = sum(f.stat().st_size for f in dest.rglob("*") if f.is_file()) / 1e6
    print(f"Đã gom {size:.1f} MB vào {dest} và {z}")
    return Path(z)
